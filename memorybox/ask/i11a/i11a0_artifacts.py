"""Excel workbook and portable results bundle for P2-I11A.0.

The Quality Review sheet is the blinded surface. Model names stay off that sheet.
"""
from __future__ import annotations

import csv
import hashlib
import json
import shutil
from pathlib import Path
from typing import Any

from memorybox.ask.i11a.i11a0_benchmark import I11A0_VERSION
from memorybox.ask.i11a.i11a0_prompt import PROMPT_VERSION, prompt_sha256

RUN_COLUMNS = (
    "run_id",
    "model",
    "digest",
    "warm_or_cold",
    "repetition",
    "confirmation",
    "thinking_mode",
    "requested_evidence_tokens",
    "evidence_sha256",
    "prompt_sha256",
    "num_ctx",
    "reserved_output_tokens",
    "classification",
    "elapsed_seconds",
    "prompt_tokens_per_second",
    "generation_tokens_per_second",
    "prompt_eval_count",
    "peak_vram_gb",
    "gpu_resident",
    "narration_relpath",
    "record_relpath",
)

QUALITY_COLUMNS = (
    "blind_id",
    "comparison_view",
    "packet_sha256",
    "narration_relpath",
    "factual_accuracy",
    "chronological_accuracy",
    "unsupported_assertions",
    "material_omissions",
    "uncertainty",
    "attribution",
    "citation",
    "coherence",
    "readability",
    "storytelling",
    "emotional_appropriateness",
    "repetition",
    "usefulness",
    "notes",
    "decision",
)

WEIGHTS = {"quality": 0.60, "evidence_capacity": 0.25, "runtime_stability": 0.15}


def _workbook():
    from openpyxl import Workbook, load_workbook

    return Workbook, load_workbook


def append_run_workbook(path: Path | str, row: dict[str, Any]) -> None:
    """Append one Runs row and save. Existing rows are kept."""
    Workbook, load_workbook = _workbook()
    target = Path(path)
    if target.is_file():
        book = load_workbook(target)
    else:
        book = Workbook()
        runs = book.active
        runs.title = "Runs"
        runs.append(list(RUN_COLUMNS))
        quality = book.create_sheet("Quality Review")
        quality.append(
            [
                "Blinded founder sheet. Score this sheet before opening sealed/identity_map.json or using the Runs sheet to identify a model."
            ]
        )
        quality.append(list(QUALITY_COLUMNS))
        config = book.create_sheet("Configuration")
        config.append(["key", "value"])
        summary = book.create_sheet("Summary")
        summary.append(["section", "item", "value"])
        summary.append(["weights", "quality", WEIGHTS["quality"]])
        summary.append(["weights", "evidence_capacity", WEIGHTS["evidence_capacity"]])
        summary.append(["weights", "runtime_stability", WEIGHTS["runtime_stability"]])
        summary.append(["comparison", "equal_input", "required"])
        summary.append(["comparison", "best_safe_setting", "required"])
        summary.append(
            [
                "decision",
                "weighted_score",
                "0.60*quality + 0.25*capacity + 0.15*stability; founder override recorded separately",
            ]
        )
    runs = book["Runs"]
    values = [row.get(column) for column in RUN_COLUMNS]
    runs.append(values)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)


def write_configuration_sheet(path: Path | str, rows: list[tuple[str, str]]) -> None:
    _Workbook, load_workbook = _workbook()
    book = load_workbook(path)
    sheet = book["Configuration"]
    for key, value in rows:
        sheet.append([key, value])
    book.save(path)


def write_summary_rows(path: Path | str, rows: list[tuple[str, str, Any]]) -> None:
    _Workbook, load_workbook = _workbook()
    book = load_workbook(path)
    sheet = book["Summary"]
    for section, item, value in rows:
        sheet.append([section, item, value])
    book.save(path)


def write_quality_row(path: Path | str, row: dict[str, Any]) -> None:
    if any(key in row for key in ("model", "digest", "tag")):
        raise ValueError("Quality Review rows must not carry model identity")
    _Workbook, load_workbook = _workbook()
    book = load_workbook(path)
    sheet = book["Quality Review"]
    sheet.append([row.get(column) for column in QUALITY_COLUMNS])
    book.save(path)


def write_results_bundle(
    *,
    bundle_dir: Path | str,
    workbook_path: Path | str,
    summary_markdown: str,
    run_rows: list[dict[str, Any]],
    narrations: dict[str, str],
    raw_responses: dict[str, str],
    telemetry: dict[str, str],
    console_logs: dict[str, str],
    validation: dict[str, str],
    config_text: str,
    prompt_text: str,
    manifests: dict[str, str],
    identity_map: dict[str, str],
) -> dict[str, Any]:
    """Copy a movable bundle. Workbook links are relative to the bundle root."""
    root = Path(bundle_dir)
    if root.exists():
        raise FileExistsError(f"bundle already exists: {root}")
    root.mkdir(parents=True)
    shutil.copyfile(workbook_path, root / "benchmark_results.xlsx")
    (root / "benchmark_summary.md").write_text(summary_markdown, encoding="utf-8", newline="\n")
    _write_table(root / "runs.csv", run_rows)
    with (root / "runs.jsonl").open("w", encoding="utf-8", newline="\n") as handle:
        for row in run_rows:
            handle.write(json.dumps(row, sort_keys=True) + "\n")
    _write_map(root / "narrations", narrations, ".txt")
    _write_map(root / "raw", raw_responses, ".json")
    _write_map(root / "telemetry", telemetry, ".jsonl")
    _write_map(root / "console", console_logs, ".log")
    _write_map(root / "validation", validation, ".json")
    (root / "config" / "benchmark_config.json").parent.mkdir(parents=True, exist_ok=True)
    (root / "config" / "benchmark_config.json").write_text(config_text, encoding="utf-8", newline="\n")
    prompt_dir = root / "prompts"
    prompt_dir.mkdir(parents=True, exist_ok=True)
    (prompt_dir / f"{PROMPT_VERSION}.txt").write_text(prompt_text, encoding="utf-8", newline="\n")
    (prompt_dir / "PROMPT_SHA256.txt").write_text(prompt_sha256() + "\n", encoding="utf-8", newline="\n")
    _write_map(root / "manifests", manifests, ".json")
    sealed = root / "sealed"
    sealed.mkdir(parents=True, exist_ok=True)
    (sealed / "identity_map.json").write_text(
        json.dumps(identity_map, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    hashes = _hash_tree(root)
    (root / "HASHES.txt").write_text(
        "\n".join(f"{digest}  {name}" for name, digest in hashes) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    return {
        "ok": True,
        "bundle": str(root),
        "version": I11A0_VERSION,
        "file_count": len(hashes),
        "models_called": False,
    }


def _write_table(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(RUN_COLUMNS)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key) for key in fieldnames})


def _write_map(directory: Path, rows: dict[str, str], suffix: str) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    for name, text in rows.items():
        safe = name.replace("/", "_").replace("\\", "_")
        (directory / f"{safe}{suffix}").write_text(text, encoding="utf-8", newline="\n")


def _hash_tree(root: Path) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for path in sorted(item for item in root.rglob("*") if item.is_file()):
        if path.name == "HASHES.txt":
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        rows.append((path.relative_to(root).as_posix(), digest))
    return rows
